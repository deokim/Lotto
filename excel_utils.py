from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def basic_sht_Style(
    worksheet, 
    font_name='나눔고딕', 
    font_size=8, 
    col_width=4.7, 
    header_height=35.0, 
    freeze_cell='B2'
):
    # 1. 전체 셀 폰트 및 사이즈 지정
    for col in range(1, worksheet.max_column + 1):
        for row in range(1, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col)
            # 첫 번째 열이나 첫 번째 행은 볼드체 적용
            if col == 1 or row == 1:
                cell.font = Font(name=font_name, size=font_size, bold=True)
            else:
                cell.font = Font(name=font_name, size=font_size)
        
        # 첫 번째 열을 제외하고 지정한 열 너비 적용
        if col == 1:
            continue
        worksheet.column_dimensions[get_column_letter(col)].width = col_width

    # 2. 헤더행(1행) 높이 및 정렬 서식 지정
    worksheet.row_dimensions[1].height = header_height
    for cell in worksheet[1]:
        cell.alignment = Alignment(wrap_text=True, vertical='top')

    # 3. 틀 고정(Freeze Pane) 적용
    worksheet.freeze_panes = freeze_cell

    #4. 화면 비율 조정
    worksheet.sheet_view.zoomScale=80

    return worksheet